"""Offline dependency smoke checks. Does not import AuditGuard or call an LLM."""
import importlib.metadata as metadata
import io
import ipaddress
import json
import os
from pathlib import Path
import platform
import socket
import sys
import tempfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
os.environ["FASTEMBED_CACHE_PATH"] = str(ROOT / "model-cache")
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["HF_HUB_DISABLE_TELEMETRY"] = "1"
RESULTS = []


def check(name, action):
    try:
        detail = action()
        RESULTS.append({"check": name, "ok": True, "detail": detail})
    except Exception as exc:
        RESULTS.append({"check": name, "ok": False, "detail": f"{type(exc).__name__}: {exc}"})


def runtime():
    assert sys.platform == "win32" and platform.machine().upper() == "AMD64"
    assert sys.version_info[:2] == (3, 13), "Requires standard CPython 3.13 x64."
    return sys.version.split()[0]


def dependencies():
    mismatches = []
    count = 0
    for line in (ROOT / "requirements.windows-py313.lock.txt").read_text().splitlines():
        if not line.strip() or line.startswith("#"):
            continue
        name, expected = line.split("==", 1)
        actual = metadata.version(name)
        count += 1
        if actual != expected:
            mismatches.append(f"{name}: {actual}, expected {expected}")
    assert not mismatches, "; ".join(mismatches)
    return f"{count} package versions match the captured environment"


def http_stack():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    app = FastAPI()
    @app.get("/ready")
    def ready():
        return {"ready": True}
    with TestClient(app) as client:
        response = client.get("/ready")
        assert response.status_code == 200 and response.json()["ready"]
    return "FastAPI + HTTPX TestClient, entirely in process"


def pdf_stack():
    from pypdf import PdfWriter, PdfReader
    import pypdfium2 as pdfium
    content = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(content)
    assert len(PdfReader(io.BytesIO(content.getvalue())).pages) == 1
    with pdfium.PdfDocument(content.getvalue()) as pdf:
        page = pdf[0]
        bitmap = page.render(scale=1)
        picture = bitmap.to_pil()
        output = io.BytesIO()
        picture.save(output, format="WEBP", lossless=True)
        assert len(output.getvalue()) > 0
        picture.close()
        bitmap.close()
        page.close()
    return "Synthetic PDF parsed, rasterized and encoded as WebP; OCR not tested or supplied"


def office_stack():
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    word = Document()
    word.add_paragraph("Setup check")
    stream = io.BytesIO()
    word.save(stream)
    assert Document(io.BytesIO(stream.getvalue())).paragraphs[0].text == "Setup check"
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Setup check"
    stream = io.BytesIO()
    slides.save(stream)
    assert Presentation(io.BytesIO(stream.getvalue())).slides[0].shapes[0].text == "Setup check"
    return "DOCX and PPTX synthetic in-memory round trips"


def vectors():
    import numpy as np
    import lancedb
    from fastembed import TextEmbedding
    from rank_bm25 import BM25Okapi
    model = TextEmbedding("BAAI/bge-small-en-v1.5", cache_dir=str(ROOT / "model-cache"), local_files_only=True, threads=2, providers=["CPUExecutionProvider"])
    docs = ["Keep authorization records for seven years.", "The cafeteria serves lunch.", "The building has two elevators."]
    matrix = np.asarray(list(model.embed(docs)), dtype=np.float32)
    assert matrix.shape == (3, 384) and np.isfinite(matrix).all()
    assert np.allclose(np.linalg.norm(matrix, axis=1), 1, atol=1e-3)
    corpus = [text.lower().split() for text in docs]
    assert int(np.argmax(BM25Okapi(corpus).get_scores(["authorization"]))) == 0
    checks_dir = ROOT / ".checks"
    checks_dir.mkdir(exist_ok=True)
    db_path = Path(tempfile.mkdtemp(prefix="lance-", dir=checks_dir))
    assert db_path.resolve().is_relative_to(checks_dir.resolve())
    db = lancedb.connect(str(db_path))
    table = db.create_table("setup_probe", [{"id": i, "vector": row.tolist()} for i, row in enumerate(matrix)])
    hit = table.search(matrix[0].tolist()).limit(1).to_list()[0]
    assert hit["id"] == 0
    return "Cached 384-dimensional CPU embeddings, BM25 and LanceDB round trip"


def main():
    check("windows_python", runtime)
    check("dependency_versions", dependencies)
    # Windows asyncio uses loopback sockets internally; allow those, block remote hosts.
    original_connect = socket.socket.connect
    def local_connect(sock, address):
        host = address[0] if isinstance(address, tuple) else ""
        try:
            local = host == "localhost" or ipaddress.ip_address(host).is_loopback
        except ValueError:
            local = False
        if not local:
            raise RuntimeError("External networking disabled during preflight")
        return original_connect(sock, address)
    with patch.object(socket.socket, "connect", local_connect):
        check("http_dependencies", http_stack)
        check("pdf_rendering", pdf_stack)
        check("office_parsers", office_stack)
        check("offline_embeddings_and_search", vectors)
    passed = sum(item["ok"] for item in RESULTS)
    print(json.dumps({"passed": passed, "total": len(RESULTS), "scope": "Setup checks only. No RAG accuracy measurement, application tests, or ALMA call.", "checks": RESULTS}, indent=2))
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    raise SystemExit(main())

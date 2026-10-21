import io
import re
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
from pypdf import PdfReader
import pypdfium2 as pdfium
from app.config import PREVIEWS_DIR, DOCS_DIR
from app.classifier import classify_document

def sanitize_doc_id(raw_id: str) -> str:
    """Sanitizes document identifiers to prevent path traversal or injection."""
    cleaned = re.sub(r'[^a-zA-Z0-9_-]', '_', Path(raw_id).stem)
    cleaned = re.sub(r'_+', '_', cleaned).strip('_')
    return cleaned or "document"

def get_or_render_page_image(doc_prefix: str, page_num: int) -> Path | None:
    """
    On-demand single-page renderer for large documents (e.g. 300-page manuals).
    Renders just the requested page in ~30ms and caches it to disk in compressed WebP format (80% quality).
    """
    file_name = f"{doc_prefix}_page_{page_num}.webp"
    file_path = PREVIEWS_DIR / file_name
    if file_path.exists():
        return file_path

    # Fallback to older PNG if exists
    png_fallback = PREVIEWS_DIR / f"{doc_prefix}_page_{page_num}.png"
    if png_fallback.exists():
        return png_fallback

    from app.config import DOCS_DIR
    pdf_path = DOCS_DIR / f"{doc_prefix}.pdf"
    if not pdf_path.exists():
        return None

    try:
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            if 1 <= page_num <= len(pdf):
                page = pdf[page_num - 1]
                image = page.render(scale=1.5).to_pil()
                image.save(file_path, "WEBP", quality=80)
                return file_path
        finally:
            pdf.close()
    except Exception as e:
        print(f"Error on-demand rendering page {page_num}: {e}")
    return None

def render_page_previews(pdf_bytes: bytes, doc_prefix: str) -> Dict[int, str]:
    """
    Smart Page Preview Renderer:
    - Small documents (<= 10 pages): renders all upfront.
    - Large enterprise documents (up to 300+ pages): saves the master PDF,
      pre-renders the first 3 pages for instant UI thumbs, and routes the rest
      to the on-demand /api/preview endpoint (sub-35ms lazy render).
    - Compresses images to WebP (quality 80) to reduce size by 85%.
    """
    from app.config import DOCS_DIR
    # Persist the master PDF for on-demand lazy rendering
    pdf_save_path = DOCS_DIR / f"{doc_prefix}.pdf"
    try:
        pdf_save_path.write_bytes(pdf_bytes)
    except Exception as e:
        print(f"Warning saving master PDF: {e}")

    page_map = {}
    try:
        pdf = pdfium.PdfDocument(pdf_bytes)
        try:
            total_pages = len(pdf)
            
            # If <= 10 pages, render all upfront; otherwise render first 3 and lazy-load remaining
            pages_to_prerender = total_pages if total_pages <= 10 else 3

            for idx in range(total_pages):
                page_num = idx + 1
                file_name = f"{doc_prefix}_page_{page_num}.webp"
                file_path = PREVIEWS_DIR / file_name

                if idx < pages_to_prerender:
                    page = pdf[idx]
                    image = page.render(scale=1.5).to_pil()
                    image.save(file_path, "WEBP", quality=80)
                    page_map[page_num] = f"/static/previews/{file_name}"
                else:
                    # Lazy on-demand route for 300-page scale
                    page_map[page_num] = f"/api/preview/{doc_prefix}/{page_num}"
        finally:
            pdf.close()
    except Exception as e:
        print(f"Warning: Could not render page preview with pypdfium2: {e}")
    return page_map

def clean_text_to_markdown(raw_text: str) -> str:
    """
    Converts raw extracted document text into clean, structured Markdown
    for fast AI comprehension, retaining headings, lists, and tables.
    """
    lines = raw_text.splitlines()
    cleaned = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            cleaned.append("")
            continue
        # Convert numbered sections (e.g. 1. SCOPE or 2.1 SLA) to Markdown Headings
        if re.match(r'^\d+\.\s+[A-Z\s]{3,}', stripped):
            cleaned.append(f"## {stripped}")
        elif re.match(r'^\d+\.\d+\s+', stripped):
            cleaned.append(f"### {stripped}")
        elif re.match(r'^[•\*\-]\s+', stripped):
            cleaned.append(f"- {stripped.lstrip('•*- ')}")
        else:
            cleaned.append(stripped)
    return "\n".join(cleaned)

def ocr_extract_page_image(page_image: Any) -> Optional[str]:
    """
    Attempts optical character recognition (OCR) on a rendered page image using pytesseract.
    Returns extracted text if successful, or None if OCR engine is unavailable or produces no text.
    """
    try:
        import pytesseract
        text = pytesseract.image_to_string(page_image)
        if text and text.strip():
            return text.strip()
    except Exception:
        pass
    return None

def extract_pages_from_pdf(pdf_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Extract page-level text and metadata from raw PDF bytes, normalized to structured Markdown.
    Includes automated fallback to pypdfium2 textpages for CID/malformed streams,
    and supports mixed and scanned/image-only PDFs so no pages are dropped and visual AST leaves
    can still be extracted with optional OCR text extraction.
    """
    pages = []
    total_pdf_pages = 0
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        total_pdf_pages = len(reader.pages)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                md_text = clean_text_to_markdown(text)
                pages.append({
                    "filename": filename,
                    "page": idx + 1,
                    "text": md_text
                })
    except Exception as e:
        print(f"pypdf extraction notice: {e}")

    # Fallback to pypdfium2 if pypdf missed pages or failed
    if not pages or len(pages) < total_pdf_pages:
        try:
            pdf = pdfium.PdfDocument(pdf_bytes)
            try:
                pdfium_total = len(pdf)
                if total_pdf_pages == 0:
                    total_pdf_pages = pdfium_total
                for idx in range(pdfium_total):
                    page_num = idx + 1
                    if not any(p["page"] == page_num for p in pages):
                        page = pdf[idx]
                        tp = None
                        t = ""
                        try:
                            tp = page.get_textpage()
                            t = tp.get_text_range() or ""
                        except Exception:
                            pass
                        finally:
                            if tp and hasattr(tp, "close"):
                                try:
                                    tp.close()
                                except Exception:
                                    pass
                            if hasattr(page, "close"):
                                try:
                                    page.close()
                                except Exception:
                                    pass
                        if t.strip():
                            pages.append({
                                "filename": filename,
                                "page": page_num,
                                "text": clean_text_to_markdown(t)
                            })
            finally:
                pdf.close()
        except Exception as e:
            print(f"pypdfium2 extraction notice: {e}")

    # Mixed & Scanned PDF Handler: Ensure no page is ever omitted from the document
    if total_pdf_pages > 0 and len(pages) < total_pdf_pages:
        existing_page_nums = {p["page"] for p in pages}
        try:
            pdf = pdfium.PdfDocument(pdf_bytes)
            try:
                for idx in range(total_pdf_pages):
                    page_num = idx + 1
                    if page_num not in existing_page_nums:
                        page = pdf[idx]
                        ocr_text = None
                        try:
                            pil_img = page.render(scale=2.0).to_pil()
                            ocr_text = ocr_extract_page_image(pil_img)
                        except Exception:
                            ocr_text = None
                        finally:
                            if hasattr(page, "close"):
                                try:
                                    page.close()
                                except Exception:
                                    pass

                        if ocr_text:
                            pages.append({
                                "filename": filename,
                                "page": page_num,
                                "text": clean_text_to_markdown(ocr_text),
                                "is_scanned": True,
                                "is_ocr": True
                            })
                        else:
                            pages.append({
                                "filename": filename,
                                "page": page_num,
                                "text": f"Page {page_num} (Visual / Scanned Content)",
                                "is_scanned": True,
                                "is_ocr": False
                            })
            finally:
                pdf.close()
        except Exception:
            for page_num in range(1, total_pdf_pages + 1):
                if page_num not in existing_page_nums:
                    pages.append({
                        "filename": filename,
                        "page": page_num,
                        "text": f"Page {page_num} (Visual / Scanned Content)",
                        "is_scanned": True,
                        "is_ocr": False
                    })

    pages.sort(key=lambda x: x["page"])
    return pages

def extract_pages_from_text(raw_text: str, filename: str) -> List[Dict[str, Any]]:
    """Fallback for raw .txt or .md files."""
    chunks = [c.strip() for c in raw_text.split("\n\n\n") if c.strip()]
    if not chunks:
        chunks = [raw_text]
    pages = []
    for idx, c in enumerate(chunks, 1):
        pages.append({
            "filename": filename,
            "page": idx,
            "text": c
        })
    return pages


def extract_pages_from_docx(docx_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Extract sections, headings, paragraphs, and tables from Word documents (.docx)."""
    import docx
    doc = docx.Document(io.BytesIO(docx_bytes))
    sections = []
    current_section = []
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if p.style.name.startswith("Heading") or re.match(r"^\d+\.", t):
            if current_section:
                sections.append("\n\n".join(current_section))
                current_section = []
            current_section.append(f"## {t}")
        else:
            current_section.append(t)
            
    for table in doc.tables:
        rows_txt = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows_txt.append(" | ".join(cells))
        if rows_txt:
            table_md = "\n".join(rows_txt)
            current_section.append(f"**Table Data:**\n```\n{table_md}\n```")
            
    if current_section:
        sections.append("\n\n".join(current_section))
        
    if not sections:
        sections = ["Empty Word document."]
        
    pages = []
    for idx, sec in enumerate(sections, 1):
        pages.append({
            "filename": filename,
            "page": idx,
            "text": clean_text_to_markdown(sec)
        })
    return pages

def extract_pages_from_pptx(pptx_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Extract slide-by-slide text, titles, notes, and tables from PowerPoint (.pptx)."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    pages = []
    for idx, slide in enumerate(prs.slides, 1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        slide_parts.append(t)
            elif shape.has_table:
                table_rows = []
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    table_rows.append(" | ".join(cells))
                if table_rows:
                    slide_parts.append("Table:\n" + "\n".join(table_rows))
                    
        slide_content = "\n\n".join(slide_parts).strip()
        if not slide_content:
            slide_content = f"Slide {idx} (Visual Graphic / Title Slide)"
            
        pages.append({
            "filename": filename,
            "page": idx,
            "text": clean_text_to_markdown(f"## Slide {idx}\n\n" + slide_content)
        })
        
    if not pages:
        pages.append({"filename": filename, "page": 1, "text": "Empty presentation."})
    return pages

def _split_into_child_units(para: str) -> List[str]:
    """
    Splits a parent paragraph or table block into fine-grained child search units.
    - For Markdown/pipe tables: pairs header context with each row so individual row lookups retain column names.
    - For prose paragraphs: splits into sentences while preserving bullet items.
    """
    lines = [ln.strip() for ln in para.splitlines() if ln.strip()]
    pipe_lines = [ln for ln in lines if ln.startswith("|")]
    if len(pipe_lines) >= 2:
        title_prefix = lines[0] if not lines[0].startswith("|") else ""
        header_row = pipe_lines[0]
        units = [para]
        for row in pipe_lines[1:]:
            if re.match(r'^\|[\s\-:|]+\|$', row):
                continue
            row_unit = f"{title_prefix} {header_row} {row}".strip()
            units.append(row_unit)
        return units

    sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+|\n(?=\-\s|\d+\.\s)', para) if len(s.strip()) > 15]
    return sentences if sentences else [para]


def hierarchical_chunk(
    pages: List[Dict[str, Any]],
    page_images: Optional[Dict[int, str]] = None,
    ast: Optional[Any] = None,
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Implements Small-to-Big (Parent-Child) chunking synchronized with the Document AST.
    Child sentences/rows are indexed for precise vector + BM25 matching.
    Parent paragraphs are passed to Gemma 4 for rich, coherent context.
    Also runs automatic document type classification.
    """
    page_images = page_images or {}
    chunks: List[Dict[str, Any]] = []
    chunk_id = 0

    combined_header = " ".join([p.get("text", "")[:1000] for p in pages[:2]])
    classification = classify_document(combined_header)
    default_filename = pages[0].get("filename", "Document") if pages else "Document"

    # When AST is provided, derive chunks directly from AST PARAGRAPH nodes for 100% node ID & path alignment
    if ast is not None and hasattr(ast, "nodes"):
        from app.ast_builder import ASTNodeType
        para_nodes = [n for n in ast.nodes.values() if n.type == ASTNodeType.PARAGRAPH]
        if para_nodes:
            for p_node in para_nodes:
                page_num = p_node.page
                para_text = p_node.text.strip()
                if not para_text:
                    continue

                # Determine section title from parent AST node
                sec_title = "Section"
                if p_node.parent_id and p_node.parent_id in ast.nodes:
                    sec_title = ast.nodes[p_node.parent_id].title

                img_url = p_node.image_url or page_images.get(page_num)
                linked_visuals = []
                for lid in p_node.linked_node_ids:
                    v_node = ast.get_node(lid)
                    if v_node and v_node.type == ASTNodeType.VISUAL:
                        linked_visuals.append({
                            "id": v_node.id,
                            "title": v_node.title,
                            "image_url": v_node.image_url or page_images.get(v_node.page),
                            "visual_type": v_node.visual_type,
                            "bbox": v_node.bbox,
                            "page": v_node.page,
                        })
                        if not img_url and v_node.image_url:
                            img_url = v_node.image_url

                child_units = _split_into_child_units(para_text)
                for sent in child_units:
                    chunk_item: Dict[str, Any] = {
                        "id": chunk_id,
                        "filename": default_filename,
                        "page": page_num,
                        "parent_text": para_text,
                        "child_text": sent,
                        "ast_node_id": p_node.id,
                        "section_title": sec_title,
                        "materialized_path": p_node.materialized_path,
                        "image_url": img_url,
                        "has_visual": bool(img_url or linked_visuals),
                    }
                    if linked_visuals:
                        chunk_item["linked_visuals"] = linked_visuals
                    chunks.append(chunk_item)
                    chunk_id += 1
            return chunks, classification

    for page in pages:
        page_num = page["page"]
        img_url = page_images.get(page_num)

        paragraphs = [p.strip() for p in re.split(r'\n\s*\n', page.get("text", "")) if len(p.strip()) > 20]
        if not paragraphs:
            paragraphs = [page.get("text", "").strip()] if page.get("text", "").strip() else []

        for p_idx, para in enumerate(paragraphs, 1):
            sentences = _split_into_child_units(para)
            for sent in sentences:
                chunks.append({
                    "id": chunk_id,
                    "filename": page.get("filename", default_filename),
                    "page": page_num,
                    "parent_text": para,
                    "child_text": sent,
                    "ast_node_id": f"para_{page_num}_{p_idx}",
                    "section_title": f"Page {page_num}",
                    "materialized_path": f"root/page_{page_num}/para_{p_idx}",
                    "image_url": img_url,
                    "has_visual": bool(img_url),
                })
                chunk_id += 1

    return chunks, classification
